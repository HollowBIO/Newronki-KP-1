"""Builds the KP2 report from the template. Run after all training/plot scripts.

Rebuilding overwrites the .docx: work manual Word edits into this script, not the file.
"""
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

REPORT_DIR = Path(__file__).resolve().parent
TEMPLATE = REPORT_DIR / "KP2_report_template.docx"
KP2 = REPORT_DIR.parent
OUT = REPORT_DIR / "ІК-61мп Балабан Ілля Олегович - KP2.docx"
RESULTS = KP2 / "Results"

if OUT.exists():
    shutil.copy2(OUT, OUT.with_name(OUT.stem + ".backup.docx"))

doc = Document(TEMPLATE)


def find(prefix, nth=None):
    """The template prompt starting with prefix; nth picks one when the prompt repeats."""
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    if nth is None:
        assert len(matches) == 1, f"{prefix!r}: {len(matches)} matches"
        nth = 0
    return matches[nth]


def fill(p, text=None):
    """Keep the bold label runs of a template prompt, replace the instruction with text."""
    for r in list(p.runs):
        if not r.bold:
            r._r.getparent().remove(r._r)
    if text:
        if not p.runs[-1].text.endswith(" "):
            p.runs[-1].text += " "
        p.add_run(text)
    return p


class Cursor:
    """Inserts new blocks one after another, starting right after an anchor paragraph."""

    def __init__(self, anchor):
        self.el = anchor._p
        self.parent = anchor._parent

    def _new_para(self, align):
        new = OxmlElement("w:p")
        self.el.addnext(new)
        self.el = new
        p = Paragraph(new, self.parent)
        p.alignment = align
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        return p

    def text(self, text, label=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None, keep_next=False):
        p = self._new_para(align)
        if label:
            p.add_run(label + " ").bold = True
        run = p.add_run(text)
        if size:
            run.font.size = Pt(size)
        p.paragraph_format.keep_with_next = keep_next

    def code(self, lines):
        for i, line in enumerate(lines):
            p = self._new_para(WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.keep_with_next = i < len(lines) - 1
            run = p.add_run(line)
            run.font.name = "Courier New"
            run.font.size = Pt(10.5)

    def picture(self, path, width_cm, caption):
        p = self._new_para(WD_ALIGN_PARAGRAPH.CENTER)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=Cm(width_cm))
        self.text(caption, align=WD_ALIGN_PARAGRAPH.CENTER, size=12)

    def table(self, caption, rows, widths_cm, size=12):
        self.text(caption, size=12, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.width = Cm(widths_cm[c])
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.0
                para.paragraph_format.keep_with_next = r < len(rows) - 1  # keep the table on one page
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.CENTER
                run = para.add_run(str(value))
                run.font.size = Pt(size)
                run.bold = r == 0
        self.el.addnext(table._tbl)
        self.el = table._tbl


# --- Header -------------------------------------------------------------------------
fill(find("ПІБ:"), "Балабан Ілля Олегович")
fill(find("Група:"), "ІК-61мп")
fill(find("Номер варіанту:"), "2")
fill(find("Завдання для варіанту:"), "датасет Rock_paper_scissors; вплив кількості та порядку шарів "
                                     "згортки та пулінгу в базовій CNN.")
fill(find("Засоби виконання практикуму:"),
     "Windows 11, CPU AMD Ryzen 5 5600 (без GPU); Python 3.12.3; TensorFlow 2.21.0 / Keras 3.15.1; "
     "NumPy, pandas, openpyxl – обробка та збереження результатів; scikit-learn – метрики; "
     "Matplotlib – графіки; Pillow – перевірка зображень.")

# --- Dataset --------------------------------------------------------------------------
p = fill(find("Набір даних (датасет):"),
         "Rock Paper Scissors (TensorFlow Datasets, автор L. Moroney) – згенеровані комп’ютерною "
         "графікою зображення кисті з жестом paper, rock або scissors; багатокласова класифікація на "
         "3 класи. Навчальна вибірка – 2520 зображень (по 840 на клас), тестова – 372 (по 124); класи "
         "збалансовані. Формат – PNG 300×300 RGBA на світлому фоні.")
Cursor(p).text("навчальні зображення – це 21 модель руки (по 7 на клас) × 120 кадрів із поступовим "
               "поворотом, тож сусідні кадри майже однакові. Тестова вибірка містить інші руки "
               "(12 × 31 кадр), помітно відмінні від навчальних: кулак анфас (у навчальній – збоку), "
               "манікюр, менший масштаб руки в кадрі. Тому тест суворо перевіряє узагальнення.",
               label="Особливості:")

# --- Preprocessing ----------------------------------------------------------------------
p = fill(find("Попередня обробка даних:"),
         "перевірка даних (окремий скрипт на кожен крок) проблем не виявила: пошкоджених файлів і "
         "дублікатів (зокрема між навчальною та тестовою вибірками) немає, усі зображення 300×300 RGBA "
         "без прозорих пікселів, мітки коректні, значення пікселів 3–255.")
cur = Cursor(p)
cur.text("1) RGBA → RGB (альфа-канал не містить інформації); 2) зменшення до 150×150 – у 4 рази "
         "менше обчислень на CPU, форма жесту зберігається; 3) нормалізація до [0, 1] шаром "
         "Rescaling(1/255) у моделі; 4) випадковий поділ навчальної вибірки 80/20 окремо в кожному "
         "класі (seed = 42, табл. 1); тестова вибірка – лише для фінальної оцінки. Через випадковий "
         "поділ валідація містить кадри тих самих рук, що й навчальна вибірка, тому вона "
         "оптимістична; реальне узагальнення показує тест.", label="Обробка:")
cur.table("Таблиця 1 – Розподіл даних за вибірками", [
    ["Вибірка", "Кількість зображень", "На клас", "Частка"],
    ["Навчальна", "2016", "672", "69,7 %"],
    ["Валідаційна", "504", "168", "17,4 %"],
    ["Тестова", "372", "124", "12,9 %"],
], [4.5, 4.5, 3.5, 3.5])

# --- Model ------------------------------------------------------------------------------
p = fill(find("Модель машинного навчання:"),
         "три блоки «згортка + пулінг» (Conv2D 3×3, ReLU; MaxPooling2D 2×2) з 32, 64 і 128 фільтрами, "
         "далі Flatten, Dense(128, ReLU) і Dense(3, softmax). Вхід – 150×150×3, вихід – імовірності "
         "класів paper/rock/scissors. Розмірності: 150×150×3 → 74×74×32 → 36×36×64 → 17×17×128 → "
         "36 992 → 128 → 3. З 4 828 739 параметрів 98 % припадає на Dense(128) після Flatten. Dropout "
         "свідомо не використано, щоб оцінити перенавчання.")
Cursor(p).code([
    "model = keras.Sequential([",
    "    keras.layers.Input(shape=(150, 150, 3)),",
    "    keras.layers.Rescaling(1.0 / 255),",
    "    keras.layers.Conv2D(32, 3, activation=\"relu\"),",
    "    keras.layers.MaxPooling2D(),",
    "    keras.layers.Conv2D(64, 3, activation=\"relu\"),",
    "    keras.layers.MaxPooling2D(),",
    "    keras.layers.Conv2D(128, 3, activation=\"relu\"),",
    "    keras.layers.MaxPooling2D(),",
    "    keras.layers.Flatten(),",
    "    keras.layers.Dense(128, activation=\"relu\"),",
    "    keras.layers.Dense(3, activation=\"softmax\"),",
    "])",
])

# --- Training ---------------------------------------------------------------------------
p = fill(find("Навчання моделі:"),
         "Adam (learning_rate = 0,001) – адаптивний оптимізатор, стабільний без підбору кроку; втрати – "
         "sparse_categorical_crossentropy (стандарт для softmax-класифікації з цілими мітками); метрика "
         "– accuracy (класи збалансовані), на тесті також precision, recall, F1 і матриця помилок. "
         "Батч 32, до 30 епох; EarlyStopping (val_loss, patience = 5) і ModelCheckpoint зберігають "
         "найкращу модель.")
Cursor(p).code([
    "model.compile(optimizer=keras.optimizers.Adam(learning_rate=0.001),",
    "              loss=\"sparse_categorical_crossentropy\", metrics=[\"accuracy\"])",
    "early_stop = keras.callbacks.EarlyStopping(",
    "    monitor=\"val_loss\", patience=5, restore_best_weights=True)",
    "checkpoint = keras.callbacks.ModelCheckpoint(",
    "    \"base_cnn.keras\", monitor=\"val_loss\", save_best_only=True)",
    "history = model.fit(train_ds, validation_data=val_ds, epochs=30,",
    "                    callbacks=[early_stop, checkpoint, timer])",
])

# --- Training results -------------------------------------------------------------------
p = fill(find("Результати навчання:", nth=0),  # 1st: base model, 2nd: pretrained
         "EarlyStopping зупинив навчання після 16 епох (≈ 12 с на епоху, ≈ 3,2 хв); найкраща – епоха "
         "10 (нумерація з 0). Точність 100 % на валідації досягнуто вже на епосі 1 (табл. 2, рис. 1).")
cur = Cursor(p)
cur.table("Таблиця 2 – Результати навчання та тестування базової моделі", [
    ["Показник", "Значення"],
    ["Кількість епох до зупинки", "16 (найкраща – епоха 10)"],
    ["Час навчання", "≈ 192 с (≈ 12 с на епоху)"],
    ["Навчальна вибірка: accuracy / loss", "100 % / 0,000027"],
    ["Валідаційна вибірка: accuracy / loss", "100 % / 0,00019"],
    ["Тестова вибірка: accuracy / loss", "89,8 % / 0,400"],
    ["Тест: macro precision / recall / F1", "0,902 / 0,898 / 0,899"],
], [9.0, 7.0])
cur.picture(RESULTS / "base_cnn_history.png", 16,
            "Рисунок 1 – Динаміка навчання базової моделі (loss у логарифмічному масштабі)")
cur.table("Таблиця 3 – Матриця помилок і метрики за класами на тестовій вибірці", [
    ["Справжній клас", "Прогноз paper", "Прогноз rock", "Прогноз scissors", "Precision", "Recall", "F1"],
    ["paper", "116", "0", "8", "0,841", "0,935", "0,885"],
    ["rock", "2", "114", "8", "1,000", "0,919", "0,958"],
    ["scissors", "20", "0", "104", "0,867", "0,839", "0,852"],
], [2.9, 2.3, 2.3, 2.6, 2.1, 1.9, 1.9])

# --- Evaluation -------------------------------------------------------------------------
p = fill(find("Оцінка результатів навчання:"),
         "недонавчання немає (100 % на навчальній вибірці з епохи 2), але є перенавчання. На кривих "
         "воно ледь помітне (після епохи 10 val loss зростає з 0,00019 до 0,00023), бо валідаційні "
         "кадри майже повторюють навчальні; на тесті з новими руками розрив очевидний – 100 % проти "
         "89,8 %. Модель значною мірою запам’ятала 21 навчальну руку, а не загальну форму жесту.")
cur = Cursor(p)
cur.text("36 із 38 помилок припадають на 4 тестові руки, несхожі на навчальні (табл. 3, рис. 2): "
         "ножиці з вузько розставленими пальцями → папір (20 помилок), кулак анфас → ножиці або папір, "
         "долоня з манікюром → ножиці. Частина хибних відповідей має впевненість 96–100 %; клас rock "
         "має precision 1,000.", label="Аналіз помилок:")
cur.picture(RESULTS / "base_cnn_error_examples.png", 16,
            "Рисунок 2 – Приклади помилок базової моделі на тестовій вибірці")
cur.text("аугментація (повороти, віддзеркалення, масштаб, яскравість), регуляризація (Dropout, L2), "
         "GlobalAveragePooling2D замість Flatten, поділ на навчальну й валідаційну вибірки за моделями "
         "рук. Загалом CNN добре підходить для задачі (89,8 % на нових руках без аугментації), але їй "
         "бракує різноманіття навчальних даних.", label="Можливі покращення:")

# --- Research: task, method, expectations -----------------------------------------------
p = fill(find("Задача дослідження:"), "вплив кількості та порядку шарів згортки та пулінгу в базовій CNN.")
cur = Cursor(p)
cur.text("змінювалася лише згорткова частина, решта – як у базовій моделі (Flatten → Dense(128) → "
         "Dense(3), Adam 0,001, батч 32, до 30 епох, EarlyStopping, seed = 42). Позначення: C<n> – "
         "Conv2D з n фільтрами 3×3 і ReLU, P – MaxPooling2D 2×2; блок – група шарів, що повторюється. "
         "Варіанти (табл. 4): one_block, base, five_blocks – кількість блоків «згортка + пулінг» 1, 3, 5 "
         "(більше неможливо: після п’ятого пулінгу карта має розмір 2×2); pool_first і pool_last – інший "
         "порядок тих самих 3 згорткових і 3 пулінгових шарів; double_conv_pool – 3 блоки з двох згорток "
         "і двох пулінгів. Кожен варіант навчено один раз; критерій – точність на тесті (на валідації "
         "всі мають ≈ 100 %).", label="Методика:")
cur.text("кожен пулінг удвічі збільшує рецептивне поле, тож більше блоків – більші ознаки (від "
         "контурів до пальців і кисті) й менший шар Dense після Flatten, отже, краще узагальнення. "
         "Пулінг на початку – швидше, але з втратою дрібних деталей; пулінг у кінці – значно більше "
         "обчислень і мале рецептивне поле. Два пулінги поспіль у кожному блоці зменшують ширину й "
         "висоту карт ознак у 4 рази (у першому блоці – з 146×146 до 36×36), тому рецептивне поле "
         "охоплює все зображення, а в Flatten надходить мінімум значень, проте є ризик утрати деталей "
         "між блоками.", label="Теоретичні очікування:")
cur.table("Таблиця 4 – Досліджувані варіанти архітектури", [
    ["Варіант", "Шари", "Параметрів", "Рецептивне поле, пікс."],
    ["one_block", "C32-P", "22 431 107", "4"],
    ["base", "C32-P-C64-P-C128-P", "4 828 739", "22"],
    ["five_blocks", "C32-P-C64-P-C128-P-C128-P-C128-P", "454 467", "94"],
    ["pool_first", "P-C32-P-C64-P-C128", "3 780 163", "36"],
    ["pool_last", "C32-C64-C128-P-P-P", "5 402 179", "14"],
    ["double_conv_pool", "C32-C32-P-P-C64-C64-P-P-C128-C128-P-P", "303 907", "148"],
], [3.5, 8.0, 2.3, 2.7], size=11)

# --- Research: results --------------------------------------------------------------------
p = fill(find("Результати експериментального дослідження:"),
         "наведено в табл. 5 і на рис. 3–4. Повторне навчання base дало ідентичний результат (89,8 %), "
         "тобто навчання з фіксованим seed відтворюване.")
cur = Cursor(p)
cur.table("Таблиця 5 – Результати варіантів (час епохи – медіана; у дужках – найкраща епоха)", [
    ["Варіант", "Епох", "Час епохи, с", "Val accuracy", "Test accuracy", "Test loss", "Macro F1"],
    ["one_block", "26 (20)", "12,5", "99,8 %", "75,5 %", "0,971", "0,747"],
    ["base", "16 (10)", "10,9", "100 %", "89,8 %", "0,400", "0,899"],
    ["five_blocks", "30 (28)", "10,5", "100 %", "90,6 %", "0,544", "0,906"],
    ["pool_first", "30 (29)", "3,3", "100 %", "83,1 %", "0,812", "0,834"],
    ["pool_last", "9 (3)", "80,6", "100 %", "86,0 %", "0,853", "0,854"],
    ["double_conv_pool", "30 (25)", "13,0", "100 %", "91,7 %", "0,288", "0,913"],
], [3.5, 1.9, 2.0, 2.2, 2.4, 2.0, 2.0], size=11)
cur.picture(RESULTS / "research" / "research_accuracy_time.png", 16,
            "Рисунок 3 – Точність на валідації й тесті (відрізок – розрив) та час епохи")
cur.text("double_conv_pool розпізнав усі 124 зображення каменю й ножиць, але лише 93 зображення "
         "паперу; five_blocks – 122 зображення ножиць (base – 104); pool_last – лише 80 ножиць. "
         "five_blocks, pool_first і double_conv_pool дійшли до ліміту 30 епох, pool_last зупинився "
         "після 9 (рис. 4).")
cur.picture(RESULTS / "research" / "research_val_loss.png", 16,
            "Рисунок 4 – Динаміка val loss: кількість блоків, порядок шарів і склад блоку")

# --- Research: conclusions -----------------------------------------------------------------
cur = Cursor(fill(find("Висновки за результатами дослідження:")))
cur.text("точність на тесті зростає 75,5 % → 89,8 % → 90,6 % за зменшення кількості параметрів у "
         "50 разів (22,4 млн → 0,45 млн), що відповідає теорії. З одним блоком рецептивне поле – "
         "4 пікселі, і мережа запам’ятовує навчальні зображення великим шаром Dense (розрив "
         "валідація – тест 24 %); з п’ятьма поле охоплює всю кисть (94 пікс.), а малий Dense не "
         "дає запам’ятовувати. Приріст 3 → 5 блоків (+0,8 %) – у межах випадкових коливань.",
         label="Кількість блоків:")
cur.text("серед однакових наборів шарів найкраще – чергування згортки й пулінгу (base). pool_first на "
         "6,7 % гірший, але втричі швидший; pool_last на 3,8 % гірший і в 7,4 раза повільніший, що "
         "відповідає очікуванням. Неочевидно, що pool_first має більше рецептивне поле (36 проти "
         "22 пікс.), але гірше узагальнює: пулінг сирих пікселів безповоротно втрачає дрібні контури.",
         label="Порядок шарів:")
cur.text("double_conv_pool має найвищу точність (91,7 %) і найменший test loss серед усіх моделей "
         "(0,288) за найменшої кількості параметрів (0,30 млн). Після трьох блоків карти мають розмір "
         "1×1, тож у Flatten надходить лише 128 значень, а рецептивне поле охоплює все зображення. На "
         "відміну від pool_first, агресивне зменшення відбувається після згорток, тому деталі встигають "
         "бути використані. Також випробувано блоки з двох згорток і одного пулінгу (double_conv, "
         "C-C-P × 3) – лише 81,7 %: більше згорток без додаткового пулінгу лише пришвидшує "
         "запам’ятовування; 5 блоків C-C-P-P неможливі, бо вже після третього блоку карта має розмір 1×1.",
         label="Склад блоку:")
cur.text("валідація не розрізняє варіанти (найменший val loss – у double_conv_pool, five_blocks і "
         "pool_first, хоча на тесті різниця до 8,6 %); один запуск на варіант не дозволяє вважати "
         "надійною різницю ≈ 1 %.", label="Обмеження:")
cur.text("найкраще узагальнюють архітектури, у яких більшість параметрів зосереджено в згорткових шарах, "
         "а карти ознак перед Flatten малі; пулінг при цьому має йти після згорток. Найкращий варіант – "
         "double_conv_pool (91,7 %, 0,30 млн параметрів).", label="Загальний висновок:")

# --- Pretrained network (task point 3) ------------------------------------------------------
fill(find("Обрана готова мережа:"),
     "MobileNetV2 (Google, 2018) – легка згорткова мережа для мобільних пристроїв, побудована з "
     "інвертованих залишкових блоків (inverted residual) із глибинно-роздільними згортками (depthwise "
     "separable), тому має лише 2,26 млн параметрів без класифікатора. Джерело – keras.applications; "
     "переднавчена на ImageNet (1,28 млн зображень, 1000 класів) із входом 224×224. Обрана як "
     "найшвидша із запропонованих архітектур для навчання на CPU.")

fill(find("Донавчання мережі:"),
     "вихідний класифікатор на 1000 класів видалено (include_top=False); додано Rescaling до [−1, 1] "
     "(передобробка MobileNetV2), GlobalAveragePooling2D, Dropout(0,2) і Dense(3, softmax); зображення "
     "завантажено одразу в розмірі 224×224. Етап 1: уся переднавчена частина заморожена, навчаються "
     "лише 3 843 параметри нового шару; Adam (0,001), до 10 епох. Етап 2: розморожено 30 верхніх шарів "
     "MobileNetV2 (1 530 243 параметри) і продовжено навчання з у 100 разів меншим кроком – Adam "
     "(0,00001), до 10 епох, щоб не зруйнувати переднавчені ознаки. На обох етапах шари BatchNorm "
     "працюють у режимі інференсу; втрати – sparse_categorical_crossentropy, батч 32, EarlyStopping "
     "(patience = 3).")

p = fill(find("Результати навчання:", nth=1),
         "обидва етапи пройшли всі 10 епох (≈ 22 і 26 с на епоху, ≈ 8 хв загалом); точність на "
         "навчальній і валідаційній вибірках сягнула 100 % (рис. 5). На тесті донавчання погіршило "
         "результат: 84,9 % → 81,2 % (табл. 6).")
cur = Cursor(p)
cur.picture(RESULTS / "transfer" / "mobilenet_v2_history.png", 15,
            "Рисунок 5 – Динаміка навчання MobileNetV2 (пунктир – початок донавчання)")
cur.table("Таблиця 6 – MobileNetV2 у порівнянні з власними CNN (тестова вибірка)", [
    ["Модель", "Навчуваних параметрів", "Accuracy", "Loss", "Macro F1",
     "Правильно paper / rock / scissors (зі 124)"],
    ["base", "4 828 739", "89,8 %", "0,400", "0,899", "116 / 114 / 104"],
    ["double_conv_pool", "303 907", "91,7 %", "0,288", "0,913", "93 / 124 / 124"],
    ["MobileNetV2, етап 1", "3 843", "84,9 %", "0,307", "0,838", "68 / 124 / 124"],
    ["MobileNetV2, етап 2", "1 530 243", "81,2 %", "0,638", "0,795", "54 / 124 / 124"],
], [3.9, 2.7, 2.0, 1.6, 1.8, 4.5], size=11)

p = fill(find("Оцінка результатів готової моделі і висновки:"),
         "модель вдалося навчити (100 % на навчальній і валідаційній вибірках), однак на тесті вона "
         "гірша за базову модель (84,9 % проти 89,8 %) і за найкращий варіант дослідження "
         "double_conv_pool (91,7 %). Камінь і ножиці розпізнано без жодної помилки, а всі помилки – "
         "зображення паперу, переважно віднесені до каменю (68 із 70 після донавчання). У навчальній "
         "вибірці «папір» – долоня з широко розставленими пальцями, а в тестовій пальці зведені, і така "
         "компактна витягнута кисть для ознак ImageNet більше схожа на кулак (рис. 6); для руки "
         "testpaper03 після донавчання помилкові всі 31 кадр.")
cur = Cursor(p)
cur.picture(RESULTS / "transfer" / "mobilenet_v2_paper_errors.png", 16,
            "Рисунок 6 – «Папір» у навчальній вибірці (ліворуч) і тестові зображення, які MobileNetV2 "
            "після донавчання відніс до каменю")
cur.text("донавчання посилило перенавчання: 1,53 млн навчуваних параметрів підлаштувалися під 21 "
         "навчальну руку – val loss зменшився з 0,0057 до 0,0004, а точність на тесті впала на 3,7 %; "
         "випадкова валідаційна вибірка цього не виявила. Етап 1 має низький test loss (0,307, близько "
         "до найкращого 0,288), тобто помилки зосереджені в одному класі, а не випадкові.",
         label="Перенавчання:")
cur.text("переднавчена мережа в принципі підходить для задачі, але її перевага – загальні ознаки "
         "ImageNet – не компенсує систематичної відмінності тестових даних; для невеликого синтетичного "
         "датасету достатньою й кращою виявилася компактна CNN, навчена з нуля. Покращити результат "
         "можна аугментацією (масштабування, повороти) та валідацією на окремих руках.",
         label="Висновок:")

# --- Template note "Обсяг 3-5 сторінки" and the empty lines before it ------------------------
note = find("Обсяг")._p
while (prev := note.getprevious()) is not None and prev.tag.endswith("}p") and not "".join(prev.itertext()).strip():
    prev.getparent().remove(prev)
note.getparent().remove(note)

doc.save(OUT)
print("saved", OUT)
